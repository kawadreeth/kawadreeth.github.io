import sys, pathlib
sys.path.insert(0, str(pathlib.Path(__file__).parent.parent))
from update_site import SubProject, Experience, Project, Token

def test_subproject_defaults():
    sp = SubProject(title="Test", tools=[], situation="s", task="t", action=[], result=[])
    assert sp.title == "Test"

def test_experience_defaults():
    exp = Experience(company="Acme", role="Intern", dates="2024", location="LA", subprojects=[])
    assert exp.subprojects == []

def test_project_defaults():
    proj = Project(title="Foo", tags=[], situation="s", task="t", action=[], result=[])
    assert proj.title == "Foo"

def test_token_defaults():
    tok = Token(type="BULLET_ITEM", text="hello")
    assert tok.situation == ""

from docx import Document
from update_site import iter_body_elements

def test_iter_body_elements_order():
    doc = Document()
    doc.add_paragraph("first")
    doc.add_table(rows=1, cols=1)
    doc.add_paragraph("third")
    elements = list(iter_body_elements(doc))
    kinds = [k for k, _ in elements]
    # new Document() starts with one empty paragraph, then we add: para, table, para
    assert kinds.count("table") == 1
    table_idx = kinds.index("table")
    assert "paragraph" in kinds[:table_idx]
    assert "paragraph" in kinds[table_idx+1:]

def test_iter_body_elements_empty_doc():
    doc = Document()
    # new Document() has one empty paragraph by default
    elements = list(iter_body_elements(doc))
    assert all(k == "paragraph" for k, _ in elements)

from update_site import parse_star_table

def _make_star_table(doc, situation, task):
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "S"
    table.rows[0].cells[1].text = f"SITUATION  {situation}"
    table.rows[1].cells[0].text = "T"
    table.rows[1].cells[1].text = f"TASK  {task}"
    return table

def test_parse_star_table_extracts_fields():
    doc = Document()
    table = _make_star_table(doc, "The initial problem.", "Fix it completely.")
    situation, task = parse_star_table(table)
    assert situation == "The initial problem."
    assert task == "Fix it completely."

def test_parse_star_table_strips_prefix():
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "S"
    table.rows[0].cells[1].text = "SITUATION No prefix needed."
    table.rows[1].cells[0].text = "T"
    table.rows[1].cells[1].text = "TASK Do the thing."
    situation, task = parse_star_table(table)
    assert not situation.startswith("SITUATION")
    assert not task.startswith("TASK")

def test_parse_star_table_wrong_labels_returns_none():
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "X"
    table.rows[0].cells[1].text = "Something"
    table.rows[1].cells[0].text = "Y"
    table.rows[1].cells[1].text = "Other"
    situation, task = parse_star_table(table)
    assert situation is None
    assert task is None

from update_site import tokenize_doc, KNOWN_COMPANIES

def test_tokenize_company_header():
    doc = Document()
    doc.add_paragraph("GrayMatter Robotics — Robotics Systems & Applications Intern")
    tokens = tokenize_doc(doc)
    assert tokens[0].type == "COMPANY_HEADER"
    assert "GrayMatter" in tokens[0].text

def test_tokenize_date_line():
    doc = Document()
    doc.add_paragraph("Jan 2026 – May 2026  |  surface finishing  |  Torrance, CA")
    tokens = tokenize_doc(doc)
    assert tokens[0].type == "DATE_LINE"

def test_tokenize_tools_line():
    doc = Document()
    doc.add_paragraph("SolidWorks \xb7 MeshLab \xb7 ROS2 \xb7 Python")
    tokens = tokenize_doc(doc)
    assert tokens[0].type == "TOOLS_LINE"

def test_tokenize_star_table():
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "S"
    table.rows[0].cells[1].text = "SITUATION The problem."
    table.rows[1].cells[0].text = "T"
    table.rows[1].cells[1].text = "TASK The work."
    tokens = tokenize_doc(doc)
    star_tokens = [t for t in tokens if t.type == "STAR_TABLE"]
    assert len(star_tokens) == 1
    assert star_tokens[0].situation == "The problem."
    assert star_tokens[0].task == "The work."

def test_tokenize_action_result_headers():
    doc = Document()
    doc.add_paragraph("ACTION")
    doc.add_paragraph("RESULT")
    tokens = tokenize_doc(doc)
    types = [t.type for t in tokens]
    assert "ACTION_HEADER" in types
    assert "RESULT_HEADER" in types

from update_site import parse_experiences

def _make_flat_experience_tokens():
    """Tokens for a single-subproject (flat star) experience."""
    return [
        Token(type="COMPANY_HEADER", text="TuTr Hyperloop — Mechanical Engineer Intern"),
        Token(type="DATE_LINE", text="Jun 2024 – Jul 2024  |  Hyperloop chassis  |  LA, CA"),
        Token(type="TOOLS_LINE", text="ANSYS · Siemens NX · Structural Analysis"),
        Token(type="STAR_TABLE", situation="Pod was too heavy.", task="Reduce weight by 30%."),
        Token(type="ACTION_HEADER"),
        Token(type="BULLET_ITEM", text="Conducted 1D structural analysis."),
        Token(type="RESULT_HEADER"),
        Token(type="BULLET_ITEM", text="Achieved 30% weight reduction."),
    ]

def _make_multi_subproject_tokens():
    """Tokens for an experience with two named subprojects."""
    return [
        Token(type="COMPANY_HEADER", text="Lumindt Labs — Mechanical Engineering Intern"),
        Token(type="DATE_LINE", text="Jun 2025 – Aug 2025  |  Thermal startup  |  SF, CA"),
        Token(type="SUBPROJECT_HEADER", text="Project 1: Hot-Wire System"),
        Token(type="TOOLS_LINE", text="Python · Raspberry Pi · SolidWorks"),
        Token(type="STAR_TABLE", situation="No measurement tool.", task="Build one."),
        Token(type="ACTION_HEADER"),
        Token(type="BULLET_ITEM", text="Built the sensor."),
        Token(type="RESULT_HEADER"),
        Token(type="BULLET_ITEM", text="10% uncertainty achieved."),
        Token(type="SUBPROJECT_HEADER", text="Project 2: Structural Design"),
        Token(type="TOOLS_LINE", text="SolidWorks Weldments · ASME"),
        Token(type="STAR_TABLE", situation="Heavy frame.", task="Lighten it."),
        Token(type="ACTION_HEADER"),
        Token(type="BULLET_ITEM", text="Designed in weldments."),
        Token(type="RESULT_HEADER"),
        Token(type="BULLET_ITEM", text="Safety margins met."),
    ]

def test_parse_flat_experience():
    tokens = _make_flat_experience_tokens()
    exps = parse_experiences(tokens)
    assert len(exps) == 1
    exp = exps[0]
    assert exp.company == "TuTr Hyperloop"
    assert exp.role == "Mechanical Engineer Intern"
    assert "Jun 2024" in exp.dates
    assert "LA" in exp.location
    assert len(exp.subprojects) == 1
    sp = exp.subprojects[0]
    assert sp.title == ""
    assert sp.situation == "Pod was too heavy."
    assert sp.task == "Reduce weight by 30%."
    assert sp.action == ["Conducted 1D structural analysis."]
    assert sp.result == ["Achieved 30% weight reduction."]

def test_parse_multi_subproject_experience():
    tokens = _make_multi_subproject_tokens()
    exps = parse_experiences(tokens)
    assert len(exps) == 1
    exp = exps[0]
    assert exp.company == "Lumindt Labs"
    assert len(exp.subprojects) == 2
    assert exp.subprojects[0].title == "Hot-Wire System"
    assert exp.subprojects[1].title == "Structural Design"
    assert exp.subprojects[0].action == ["Built the sensor."]
    assert exp.subprojects[1].result == ["Safety margins met."]

def test_parse_two_experiences():
    tokens = _make_flat_experience_tokens() + _make_flat_experience_tokens()
    tokens2 = list(tokens)
    tokens2[len(_make_flat_experience_tokens())] = Token(
        type="COMPANY_HEADER", text="National Institute of Wind Energy — Wind Blade Intern"
    )
    exps = parse_experiences(tokens2)
    assert len(exps) == 2
